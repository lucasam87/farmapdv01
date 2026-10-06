# Gera as fixtures de planilha usadas pelos testes de lib/xlsx.ts.
#
# NAO e dependencia do app e nao roda no build: existe para as fixtures serem
# reproduziveis (qualquer pessoa regera os mesmos arquivos).
#
# Rodar com o Python empacotado do harness, que e o unico que tem openpyxl:
#   C:\Users\lucas\.dsh\dsh-runtimes\dsh-primary-runtime\dependencies\python\python.exe ^
#       scripts\gerar-fixtures-xlsx.py
#
# Gera tres arquivos em tests/fixtures/:
#   produtos.xlsx       -- 3 produtos + 1 linha com preco em TEXTO, preco numerico
#   produtos_texto.xlsx -- precos em texto, coluna extra, ordem diferente e rodape TOTAL
#   produtos.csv        -- o MESMO conteudo de produtos.xlsx, em CSV de verdade

import csv
import os

from openpyxl import Workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(RAIZ, "tests", "fixtures")
os.makedirs(DESTINO, exist_ok=True)

CABECALHO = [
    "codigo_barras", "nome", "categoria", "unidade",
    "preco_custo", "preco_venda", "estoque_atual", "estoque_minimo",
]

# (codigo, nome, categoria, unidade, custo, venda, estoque, minimo)
# O nome tem acento de proposito: prova que o XML sai em UTF-8 e nao quebra.
PRODUTOS = [
    ("0012345", "DIPIRONA 500mg ção", "Analgésico", "un", 2.50, 4.50, 10, 5),
    ("0098765", "PARACETAMOL 750mg", "Analgésico", "un", 3.10, 6.90, 0, 2),
    ("0055555", "DIPIRONA GOTAS 20ml", "Analgésico", "ml", 5.00, 9.99, 3, 1),
]


def gravar_produtos():
    """Precos como NUMERO e codigo como TEXTO; a ultima linha tem preco em texto."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(CABECALHO)

    for codigo, nome, categoria, unidade, custo, venda, estoque, minimo in PRODUTOS:
        ws.append([codigo, nome, categoria, unidade, custo, venda, estoque, minimo])

    # Linha extra com o preco como TEXTO ("4,50"): o exportor mistura os dois
    # formatos e o leitor tem que aguentar os dois no mesmo arquivo.
    ws.append(["0077777", "PRODUTO PRECO TEXTO", "Analgésico", "un",
               "2,00", "4,50", 1, 0])

    # O codigo de barras precisa continuar texto, senao o Excel come os zeros.
    for linha in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for celula in linha:
            celula.number_format = "@"

    caminho = os.path.join(DESTINO, "produtos.xlsx")
    wb.save(caminho)
    return caminho


def gravar_produtos_texto():
    """Precos todos em texto, coluna extra no meio, ordem diferente e rodape TOTAL."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"

    # Ordem diferente de proposito: o mapeamento e por cabecalho, nao por posicao.
    ws.append(["nome", "codigo_barras", "observacao", "preco_venda", "estoque_atual"])

    linhas = [
        ("DIPIRONA 500mg ção", "0012345", "nao usar", "4,50", 10),
        ("PARACETAMOL 750mg", "0098765", "", "6,90", 0),
        ("DIPIRONA GOTAS 20ml", "0055555", "x", "9,99", 3),
    ]
    for nome, codigo, obs, preco, estoque in linhas:
        ws.append([nome, codigo, obs, preco, estoque])

    # Rodape de total: o exportor de sistema costuma terminar assim. As outras
    # celulas ficam vazias, menos a coluna de preco com a soma.
    ws.append(["TOTAL", "", "", "21,39", ""])

    caminho = os.path.join(DESTINO, "produtos_texto.xlsx")
    wb.save(caminho)
    return caminho


def gravar_csv():
    """Mesmo conteudo de produtos.xlsx, em CSV de verdade (virgula + aspas)."""
    caminho = os.path.join(DESTINO, "produtos.csv")
    with open(caminho, "w", encoding="utf-8", newline="") as arquivo:
        escritor = csv.writer(arquivo)
        escritor.writerow(CABECALHO)
        for codigo, nome, categoria, unidade, custo, venda, estoque, minimo in PRODUTOS:
            # O preco sai na forma que o Excel GUARDA o numero (2.5, e nao 2.50):
            # assim o CSV e o XLSX produzem literalmente a mesma Tabela depois de
            # lidos, e a comparacao dos dois leitores prova o que promete.
            escritor.writerow([codigo, nome, categoria, unidade,
                               f"{custo:g}", f"{venda:g}", estoque, minimo])
        escritor.writerow(["0077777", "PRODUTO PRECO TEXTO", "Analgésico", "un",
                           "2,00", "4,50", 1, 0])
    return caminho


if __name__ == "__main__":
    for caminho in (gravar_produtos(), gravar_produtos_texto(), gravar_csv()):
        print(f"gerado: {os.path.relpath(caminho, RAIZ)} ({os.path.getsize(caminho)} bytes)")
